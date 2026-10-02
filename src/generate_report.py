import pandas as pd
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

def add_heading(doc, text, level):
    heading = doc.add_heading(text, level=level)
    return heading

def add_paragraph(doc, text, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph(text)
    p.alignment = align
    return p

def main():
    doc = Document()

    # --- 1. Halaman Judul ---
    doc.add_heading('UTS Praktikum Pembelajaran Mesin', 0).alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_heading('Implementasi Decision Tree pada Dataset Online Shoppers Purchasing Intention dengan Penanganan Imbalanced Data', 1).alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph('\n\n\n')
    p = doc.add_paragraph()
    p.add_run('Disusun Oleh:\n').bold = True
    p.add_run('Nama: [Nama Mahasiswa]\n')
    p.add_run('NIM: [NIM Mahasiswa]\n')
    p.add_run('Kelas: [Kelas]\n')
    p.add_run('Kelompok: [Kelompok]\n')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()

    # --- 2. BAB I – Pendahuluan ---
    add_heading(doc, 'BAB I – Pendahuluan', 1)
    
    add_heading(doc, '1.1 Latar Belakang', 2)
    add_paragraph(doc, 'Dalam era e-commerce, memahami intensi pembelian pelanggan sangat krusial untuk meningkatkan konversi dan profitabilitas. Dataset Online Shoppers Purchasing Intention merekam berbagai sesi aktivitas pengunjung di suatu situs web. Salah satu tantangan utama dalam memprediksi apakah pengunjung akan melakukan pembelian (Revenue) adalah ketidakseimbangan kelas (imbalanced data), di mana mayoritas pengunjung hanya melakukan browsing tanpa membeli.')
    
    add_heading(doc, '1.2 Rumusan Masalah', 2)
    add_paragraph(doc, '1. Bagaimana cara menangani ketidakseimbangan kelas (imbalanced data) pada dataset Online Shoppers Purchasing Intention?\n2. Bagaimana perbandingan performa model klasifikasi Decision Tree jika dilatih dengan data hasil oversampling (SMOTENC) berbanding undersampling (RandomUnderSampler)?')
    
    add_heading(doc, '1.3 Tujuan Praktikum', 2)
    add_paragraph(doc, '1. Mengimplementasikan teknik preprocessing dan resampling pada dataset imbalanced yang memiliki fitur campuran numerik dan kategorikal.\n2. Mengevaluasi performa algoritma Decision Tree berdasarkan metrik akurasi, presisi, recall, dan F1-score.')

    # --- 3. BAB II – Landasan Teori ---
    add_heading(doc, 'BAB II – Landasan Teori', 1)
    
    add_heading(doc, '2.1 Dataset dan Imbalanced Data', 2)
    add_paragraph(doc, 'Dataset Online Shoppers Purchasing Intention berisi log sesi pengunjung. Imbalanced data terjadi ketika satu kelas target (misal False/Tidak membeli) mendominasi distribusi dibandingkan kelas lainnya.')
    
    add_heading(doc, '2.2 Data Preprocessing', 2)
    add_paragraph(doc, 'Data preprocessing adalah tahap membersihkan data kotor, termasuk menangani missing value (nilai yang kosong), duplikasi (data identik), dan outlier (pencilan ekstrem).')
    
    add_heading(doc, '2.3 Data Transformation', 2)
    add_paragraph(doc, 'StandardScaler menstandarisasi fitur numerik dengan mengurangkan mean dan membagi dengan standar deviasi. OneHotEncoder mengubah data kategorikal menjadi representasi biner array.')
    
    add_heading(doc, '2.4 Train-Test Split dan Stratifikasi', 2)
    add_paragraph(doc, 'Stratifikasi memastikan proporsi antar kelas pada set pelatihan (training) dan pengujian (testing) tetap konsisten dengan distribusi asli.')
    
    add_heading(doc, '2.5 SMOTENC dan RandomUnderSampler', 2)
    add_paragraph(doc, 'SMOTENC merupakan variasi SMOTE yang mampu menangani dataset berisi campuran fitur nominal (kategorikal) dan kontinu (numerik). RandomUnderSampler mengurangi ukuran kelas mayoritas secara acak hingga setara dengan kelas minoritas.')
    
    add_heading(doc, '2.6 Decision Tree dan Metrik Evaluasi', 2)
    add_paragraph(doc, 'Decision tree memecah data berdasarkan aturan berbasis fitur. Model dievaluasi menggunakan matriks kebingungan (Confusion Matrix) untuk mendapatkan Accuracy, Precision, Recall, dan F1-Score.')

    # --- 4. BAB III – Metodologi ---
    add_heading(doc, 'BAB III – Metodologi', 1)
    
    add_heading(doc, '3.1 Sumber dan Karakteristik Dataset', 2)
    add_paragraph(doc, 'Dataset berasal dari repositori UC Irvine yang dibagikan secara publik melalui Kaggle, terdiri atas 12.330 baris dan 18 kolom (14 numerik dan 4 kategorikal). Target klasifikasi adalah atribut Revenue.')
    
    add_heading(doc, '3.2 Tahapan Preprocessing dan Transformasi', 2)
    add_paragraph(doc, 'Pemeriksaan duplikat dilakukan sebelum data dipecah, di mana ditemukan 125 duplikat yang langsung dihapus. Tidak ada missing values. Data kategorikal tidak langsung disandikan secara dummy agar tidak mengacaukan resampler.')
    
    add_heading(doc, '3.3 Pembagian Data (Train-Test Split)', 2)
    add_paragraph(doc, 'Data dipisah 80% sebagai set pelatihan dan 20% pengujian secara stratifikasi.')
    
    add_heading(doc, '3.4 Resampling (SMOTENC dan Undersampler)', 2)
    add_paragraph(doc, 'Eksperimen dipecah dua. SMOTENC diaplikasikan sebelum OneHotEncoding untuk menangani resintesa nilai kategorikal secara benar. Di sisi lain, pipeline Undersampling diaplikasikan pada set pelatihan yang sama.')
    
    add_heading(doc, '3.5 Pemodelan Decision Tree', 2)
    add_paragraph(doc, 'Decision Tree dilatih pada masing-masing hasil resampling menggunakan data uji yang dijaga kerahasiaannya selama fase pelatihan.')

    # --- 5. BAB IV – Hasil dan Pembahasan ---
    add_heading(doc, 'BAB IV – Hasil dan Pembahasan', 1)
    
    add_heading(doc, '4.1 Hasil Preprocessing', 2)
    add_paragraph(doc, 'Data awal: 12.330 baris. Ditemukan 0 missing value, 125 baris duplikat (dihapus sehingga sisa 12.205). Outlier terdeteksi banyak pada kolom numerik (berdasarkan metode IQR), namun dipertahankan karena sifat Decision Tree yang kebal outlier.')
    
    add_heading(doc, '4.2 Distribusi Kelas (Sebelum dan Sesudah Resampling)', 2)
    add_paragraph(doc, 'Gambar 1 menunjukkan distribusi awal yang sangat imbalanced. Gambar 2 dan 3 mendemonstrasikan perataan kelas (balance) setelah SMOTENC (8.238 baris positif dan negatif) serta RandomUnderSampler (1.526 baris positif dan negatif).')
    
    def add_image(path, caption):
        if os.path.exists(path):
            doc.add_picture(path, width=Inches(5.0))
            p = doc.add_paragraph(caption)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            add_paragraph(doc, f'[Gambar tidak ditemukan: {path}]')

    add_image('../outputs/figures/dist_before_resampling.png', 'Gambar 1. Distribusi Kelas Sebelum Resampling')
    add_image('../outputs/figures/dist_after_smotenc.png', 'Gambar 2. Distribusi Kelas Setelah Oversampling (SMOTENC)')
    add_image('../outputs/figures/dist_after_rus.png', 'Gambar 3. Distribusi Kelas Setelah Undersampling (RandomUnderSampler)')

    add_heading(doc, '4.3 Matriks Kebingungan (Confusion Matrix)', 2)
    add_paragraph(doc, 'Berikut adalah matriks kebingungan dari kedua eksperimen saat diujikan pada data uji asli sebesar 20%.')
    
    add_image('../outputs/figures/cm_dt_smotenc.png', 'Gambar 4. Confusion Matrix Model SMOTENC')
    add_image('../outputs/figures/cm_dt_rus.png', 'Gambar 5. Confusion Matrix Model Undersampling')

    add_heading(doc, '4.4 Perbandingan Metrik Evaluasi', 2)
    add_paragraph(doc, 'Tabel 1 menyajikan rangkuman evaluasi metrik antara dua perlakuan.')
    
    # Load and add table
    if os.path.exists('../outputs/tables/evaluasi_model.csv'):
        df_metrics = pd.read_csv('../outputs/tables/evaluasi_model.csv')
        table = doc.add_table(rows=1, cols=len(df_metrics.columns))
        table.style = 'Table Grid'
        
        # Header
        hdr_cells = table.rows[0].cells
        for i, col in enumerate(df_metrics.columns):
            hdr_cells[i].text = col
        
        # Data
        for _, row in df_metrics.iterrows():
            row_cells = table.add_row().cells
            for i, val in enumerate(row):
                if isinstance(val, float):
                    row_cells[i].text = f"{val:.4f}"
                else:
                    row_cells[i].text = str(val)
        
        add_paragraph(doc, 'Tabel 1. Perbandingan Metrik Evaluasi (Sumber: Eksekusi Kode)').alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        add_paragraph(doc, '[Tabel evaluasi_model.csv tidak ditemukan]')

    add_image('../outputs/figures/perbandingan_metrik.png', 'Gambar 6. Grafik Perbandingan Metrik')

    add_heading(doc, '4.5 Pembahasan Perbedaan', 2)
    add_paragraph(doc, 'Dari Tabel 1 dan Gambar 6, Model Decision Tree dengan SMOTENC memberikan hasil F1-Score (sekitar 63.7%) dan Akurasi (sekitar 86.9%) yang lebih baik dibandingkan Undersampling. Undersampling unggul pada Recall (sekitar 81%), artinya sangat agresif dalam mendeteksi kelas positif, namun merugikan Precision sehingga banyak kesalahan prediksi kelas positif palsu. Keterbatasan pada eksperimen ini adalah tidak dilakukannya hiperparameter tuning (GridSearchCV) mendalam pada arsitektur pohon.')

    # --- 6. BAB V – Kesimpulan dan Saran ---
    add_heading(doc, 'BAB V – Kesimpulan dan Saran', 1)
    
    add_heading(doc, '5.1 Kesimpulan', 2)
    add_paragraph(doc, 'Penanganan imbalanced data dengan mengombinasikan pra-pemrosesan yang teliti, SMOTENC untuk oversampling yang tidak merusak data nominal, dan pembagian dataset yang disterilisasi dapat secara nyata menyeimbangkan kemampuan prediktif Decision Tree. SMOTENC direkomendasikan untuk proyek ini karena menyeimbangkan keseluruhan akurasi dan F1-Score tanpa banyak mengorbankan variasi alami.')
    
    add_heading(doc, '5.2 Saran', 2)
    add_paragraph(doc, 'Pengembangan selanjutnya direkomendasikan untuk meninjau algoritma klasifikasi *ensemble* (seperti Random Forest atau XGBoost) dan menyertakan optimalisasi *hyperparameter* (Hyperparameter Tuning).')

    # --- 7. Daftar Pustaka ---
    add_heading(doc, 'Daftar Pustaka', 1)
    add_paragraph(doc, '1. Sakar, C. O., Polat, S. O., Katircioglu, M., & Kusetogullari, Y. (2019). Real-time prediction of online shoppers purchasing intention using multilayer perceptron and LSTM recurrent neural networks. Neural Computing and Applications, 31(10), 6893-6908.')
    add_paragraph(doc, '2. Chawla, N. V., Bowyer, K. W., Hall, L. O., & Kegelmeyer, W. P. (2002). SMOTE: synthetic minority over-sampling technique. Journal of artificial intelligence research, 16, 321-357.')
    add_paragraph(doc, '3. Pedregosa, F., et al. (2011). Scikit-learn: Machine Learning in Python. JMLR 12, pp. 2825-2830.')
    add_paragraph(doc, '4. Lemaître, G., Nogueira, F., & Aridas, C. K. (2017). Imbalanced-learn: A Python Toolbox to Tackle the Curse of Imbalanced Datasets in Machine Learning. Journal of Machine Learning Research, 18(17), 1-5.')

    # Save Document
    doc.save('../Laporan_UTS_Pembelajaran_Mesin.docx')
    print("Laporan berhasil dibuat di root workspace.")

if __name__ == '__main__':
    main()
